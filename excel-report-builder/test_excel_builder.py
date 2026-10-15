"""Tests for Excel report builder."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import load_workbook

OUT = Path(__file__).parent / "monthly_sales_report.xlsx"

def test_workbook_exists():
    from excel_builder import build
    build()
    assert OUT.exists()

def test_sheets_exist():
    wb = load_workbook(OUT)
    assert set(wb.sheetnames) == {"Sales Data", "By Category", "Dashboard"}

def test_sales_data_row_count():
    wb = load_workbook(OUT)
    ws = wb["Sales Data"]
    assert ws.max_row == 17

def test_revenue_formula():
    wb = load_workbook(OUT)
    ws = wb["Sales Data"]
    assert ws["G2"].value == "=E2*F2"

def test_pivot_formulas():
    wb = load_workbook(OUT)
    ws = wb["By Category"]
    assert "SUMIFS" in str(ws["B2"].value)
    assert "COUNTIFS" in str(ws["E2"].value)

def test_dashboard_kpis():
    wb = load_workbook(OUT)
    ws = wb["Dashboard"]
    assert ws["B2"].value == "=SUM('Sales Data'!G:G)"
    assert ws["B3"].value == "=COUNTA('Sales Data'!A:A)-1"
    # Average order value depends on revenue and order count, never itself.
    from openpyxl.formula import Tokenizer
    references = {t.value for t in Tokenizer(ws["B4"].value).items if t.subtype == "RANGE"}
    assert references == {"B2", "B3"}

def test_chart_exists():
    wb = load_workbook(OUT)
    ws = wb["Dashboard"]
    assert len(ws._charts) == 1

if __name__ == "__main__":
    import pytest
    pytest.main([__file__, "-v"])
