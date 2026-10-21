from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from invoice_watch import build, load, parse_amount, parse_date, status

SAMPLE = Path(__file__).parent / "sample_faktury.csv"
TODAY = date(2026, 10, 21)


def test_parsing():
    assert parse_amount("12 500,00") == 12500.0
    assert parse_date("24.10.2026") == date(2026, 10, 24) == parse_date("2026-10-24")


def test_status():
    by_no = {r["cislo"]: status(r, TODAY) for r in load(SAMPLE)}
    assert by_no["2026001"] == "po splatnosti"
    assert by_no["2026003"] == "zaplaceno"
    assert by_no["2026005"] == "brzy splatné"
    assert by_no["2026007"] == "v pořádku"


def test_output(tmp_path):
    res = build(load(SAMPLE), tmp_path / "o.xlsx", TODAY)
    assert res["po_splatnosti"] == 12500 + 4200 + 7300 + 15000
    wb = load_workbook(tmp_path / "o.xlsx")
    assert wb["Dlužníci"]["A2"].value == "Novák Instalatérství"  # biggest debtor first
    assert wb["Dlužníci"].cell(wb["Dlužníci"].max_row, 2).value == res["po_splatnosti"]


def _row(due, paid=False):
    return {"splatnost": due, "zaplaceno": paid}


def test_status_boundaries():
    assert status(_row(date(2026, 10, 20)), TODAY) == "po splatnosti"
    assert status(_row(TODAY), TODAY) == "brzy splatné"
    assert status(_row(date(2026, 10, 28)), TODAY) == "brzy splatné"
    assert status(_row(date(2026, 10, 29)), TODAY) == "v pořádku"


def test_czech_amounts():
    assert parse_amount("1.234,50") == 1234.5
    assert parse_amount("12 500 Kč") == 12500.0
    assert parse_amount("99.5") == 99.5


def test_bad_input_reports_line(tmp_path):
    bad = tmp_path / "bad.csv"
    bad.write_text("cislo;odberatel;castka;splatnost;zaplaceno\n1;A;100;2026-10-01;ne\n2;B;;2026-10-01;ne\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Řádek 3"):
        load(bad)
    bad.write_text("cislo;odberatel\n1;A\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Chybí sloupce"):
        load(bad)
