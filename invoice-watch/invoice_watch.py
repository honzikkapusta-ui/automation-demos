"""Invoice watch: CSV of invoices -> Excel with overdue / due-soon highlighting and per-customer summary."""
import csv
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

RED = PatternFill("solid", start_color="F8B4B4")
ORANGE = PatternFill("solid", start_color="FDE3B0")
HEADER = Font(bold=True)
SOON_DAYS = 7
COLUMNS = ("cislo", "odberatel", "castka", "splatnost", "zaplaceno")


def parse_date(text: str) -> date:
    text = text.strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Neznámé datum: {text!r}")


def parse_amount(text: str) -> float:
    text = text.replace("Kč", "").replace(" ", "").replace("\xa0", "")
    if "," in text:  # Czech format: 1.234,50 -> thousands dot, decimal comma
        text = text.replace(".", "").replace(",", ".")
    return float(text)


def load(path: Path) -> list[dict]:
    """Columns: cislo, odberatel, castka, splatnost, zaplaceno (ano/ne)."""
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=";")
        missing = set(COLUMNS) - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"Chybí sloupce: {', '.join(sorted(missing))}")
        rows = []
        for line, r in enumerate(reader, start=2):
            try:
                rows.append({
                    "cislo": r["cislo"].strip(),
                    "odberatel": r["odberatel"].strip(),
                    "castka": parse_amount(r["castka"]),
                    "splatnost": parse_date(r["splatnost"]),
                    "zaplaceno": r["zaplaceno"].strip().lower() in ("ano", "yes", "1"),
                })
            except (ValueError, AttributeError) as e:
                raise ValueError(f"Řádek {line}: {e}") from e
    return rows


def status(row: dict, today: date) -> str:
    if row["zaplaceno"]:
        return "zaplaceno"
    if row["splatnost"] < today:
        return "po splatnosti"
    if row["splatnost"] <= today + timedelta(days=SOON_DAYS):
        return "brzy splatné"
    return "v pořádku"


def build(rows: list[dict], out: Path, today: date) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Faktury"
    ws.append(["Číslo", "Odběratel", "Částka (Kč)", "Splatnost", "Zaplaceno", "Dní po splatnosti", "Stav"])
    for c in ws[1]:
        c.font = HEADER
    totals: dict[str, float] = {}
    overdue_sum = 0.0
    for r in sorted(rows, key=lambda r: r["splatnost"]):
        st = status(r, today)
        days = (today - r["splatnost"]).days if st == "po splatnosti" else 0
        ws.append([r["cislo"], r["odberatel"], r["castka"], r["splatnost"],
                   "ano" if r["zaplaceno"] else "ne", days, st])
        fill = RED if st == "po splatnosti" else ORANGE if st == "brzy splatné" else None
        for c in ws[ws.max_row]:
            if fill:
                c.fill = fill
        ws.cell(ws.max_row, 3).number_format = "#,##0.00"
        ws.cell(ws.max_row, 4).number_format = "DD.MM.YYYY"
        if st == "po splatnosti":
            totals[r["odberatel"]] = totals.get(r["odberatel"], 0) + r["castka"]
            overdue_sum += r["castka"]
    for col, w in zip("ABCDEFG", (12, 28, 14, 12, 11, 18, 15)):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    sm = wb.create_sheet("Dlužníci")
    sm.append(["Odběratel", "Dluh po splatnosti (Kč)"])
    for c in sm[1]:
        c.font = HEADER
    for name, amount in sorted(totals.items(), key=lambda kv: -kv[1]):
        sm.append([name, amount])
        sm.cell(sm.max_row, 2).number_format = "#,##0.00"
    sm.append(["Celkem", overdue_sum])
    sm.cell(sm.max_row, 1).font = HEADER
    sm.cell(sm.max_row, 2).number_format = "#,##0.00"
    sm.column_dimensions["A"].width = 28
    sm.column_dimensions["B"].width = 24
    wb.save(out)
    return {"po_splatnosti": overdue_sum, "dluznici": len(totals)}


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "sample_faktury.csv"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent / "faktury_prehled.xlsx"
    try:
        res = build(load(src), dst, date.today())
    except (ValueError, OSError) as e:
        sys.exit(f"Chyba: {e}")
    money = f"{res['po_splatnosti']:,.0f}".replace(",", " ")
    print(f"Hotovo: {dst}. Po splatnosti {money} Kč u {res['dluznici']} odběratelů.")
